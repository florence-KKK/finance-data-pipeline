import os
import sys
import gc
import json
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime

# ==================== 全局错误捕获 ====================
def global_exception_handler(exc_type, exc_value, exc_traceback):
    print("\n" + "="*60)
    print("发生未捕获的异常：")
    print(f"类型：{exc_type}")
    print(f"值：{exc_value}")
    import traceback
    traceback.print_tb(exc_traceback)
    print("="*60)
    input("按回车键退出...")
    sys.exit(1)

sys.excepthook = global_exception_handler

# ==================== 敏感信息配置区（外发时替换为占位符）====================
PATHS = {
    'input_folder':     r'F:\...\BI待处理数据-项目4-导出BI数据校验-py',
    'field_map_file':   r'F:\...\field_map.json',
    'match_file':       r'F:\...\活动名称对应一二级分类.xlsx',
    'check_file1':      r'F:\...\普通分润核对.xlsx',
    'check_file2':      r'F:\...\收入成本校验表--python校验.xlsx',
    'output_file':      r'F:\...\核验结果.xlsx',
    'second_arch_file': r'F:\...\分公司目录-新二级部门标签.xlsx',
    'match_output':     r'F:\...\分类校验汇总.xlsx',
}

RULES = {
    'arch_replace_map': {
        '【旧分公司A】': '【新分公司A】',
        '【旧分公司B】': '【新分公司B】',
        '【旧分公司C】': '【新分公司C】',
        '【旧总部】':   '【新总部】',
    },
    'special_manager_id':  0,
    'special_second_arch': '【特殊分公司名称】',
    'invoice_exempt_category': '【开票豁免分类编码】',
    'allowed_arch_exceptions': ['【例外部门1】', '【例外部门2】'],

    'match_sheet_name': '名称对应活动类型',
    'log_sheet_name_check1_trans': '笔数和交易',
    'log_sheet_name_check1_profit': '分润核对',
    'log_sheet_name_check2': 'Sheet1',

    # 分类规则
    'classify_keywords': [
        '【关键词1】', '【关键词2】', '【关键词3】', '【关键词4】',
    ],
    'classify_target_category': '项目',
    'classify_direct_map': {
        '【特殊二级分类1】': '【对应分类1】',
        '【特殊二级分类2】': '【对应分类2】',
    },
    'classify_split_keyword': '普通分润',
    'special_cats': [
        '【特殊分类1】', '【特殊分类2】', '【特殊分类3】',
        '【特殊分类4】', '【特殊分类5】', '【特殊分类6】',
        '【特殊分类7】', '【特殊分类8】', '【特殊分类9】',
    ],
    'classify_special_map': {
        '【映射源分类】': ['【关键词1】', '【关键词2】'],
    },

    'detail_sheet_prefix': '明细',
}

# 脱敏后的字段名常量（代码全程使用）
COL_MERCHANT_NO   = '商户编号'
COL_MERCHANT_NAME = '商户名'
COL_EMP_ID        = '员工编号'
COL_EMP_NAME      = '员工姓名'
COL_AGENT_ID      = '渠道编号'
COL_AGENT_NAME    = '渠道名称'
COL_ARCH1         = '组织层级1'
COL_ARCH2         = '组织层级2'
COL_ARCH3         = '组织层级3'
COL_ARCH4         = '组织层级4'
COL_OPS_ID        = '运营经理编号'
COL_OPS_NAME      = '运营经理姓名'
COL_INDUSTRY      = '行业分类'
COL_SALES_ID      = '销售员编号'
COL_SALES_NAME    = '销售员姓名'
COL_TXN_COUNT     = '有效交易笔数'
COL_TXN_AMT       = '有效交易金额'
COL_REWARD        = '补贴金额'
COL_REWARD_BASE   = '补贴/基数'
COL_RATIO         = '分润比例'
COL_SETTLE        = '支出金额'
COL_PROFIT        = '净收益'
COL_CHANNEL       = '支付渠道'
COL_BIZ_CAT1      = '业务大类'
COL_BIZ_CAT2      = '业务小类'
COL_BILL_PERIOD   = '账期'
COL_BI_PERIOD     = '报表期'
COL_TABLE_NAME    = '数据表'
COL_INVOICE       = '发票抬头'
COL_PROJECT       = '项目名'
COL_PROJECT_TYPE  = '项目类型'
COL_PROJECT_CAT   = '项目类别'
COL_SETTLE_METHOD = '结算方式'
COL_AGENT_SETTLE  = '渠道结算标记'
COL_EMP_SETTLE    = '员工结算标记'
COL_EMP_D0        = '员工即时分润'
COL_AGENT_D0      = '渠道即时分润'

# 收入/成本列名（脱敏后）
COL_REWARD_INCOME = '补贴金额（收入）'
COL_SETTLE_COST   = '支出金额（成本）'
COL_CATEGORY      = '分类'
# ==========================================================================

with open(PATHS['field_map_file'], 'r', encoding='utf-8') as f:
    FIELD_MAP = json.load(f)

INPUT_FOLDER     = PATHS['input_folder']
MATCH_FILE       = PATHS['match_file']
MATCH_SHEET      = RULES['match_sheet_name']
CHECK_FILE1      = PATHS['check_file1']
CHECK_FILE2      = PATHS['check_file2']
OUTPUT_FILE      = PATHS['output_file']
SECOND_ARCH_FILE = PATHS['second_arch_file']
MATCH_OUTPUT     = PATHS['match_output']

SPECIAL_MANAGER_ID      = RULES['special_manager_id']
SPECIAL_SECOND_ARCH     = RULES['special_second_arch']
INVOICE_EXEMPT_CATEGORY = RULES['invoice_exempt_category']
ALLOWED_ARCH_EXCEPTIONS = RULES['allowed_arch_exceptions']
ARCH_REPLACE_MAP        = RULES['arch_replace_map']
SPECIAL_CATS            = RULES['special_cats']
CLASSIFY_KEYWORDS       = RULES['classify_keywords']
CLASSIFY_TARGET_CATEGORY = RULES['classify_target_category']
CLASSIFY_DIRECT_MAP     = RULES['classify_direct_map']
CLASSIFY_SPLIT_KEYWORD  = RULES['classify_split_keyword']
CLASSIFY_SPECIAL_MAP    = RULES['classify_special_map']
SHEET_CHECK1_TRANS      = RULES['log_sheet_name_check1_trans']
SHEET_CHECK1_PROFIT     = RULES['log_sheet_name_check1_profit']
SHEET_CHECK2            = RULES['log_sheet_name_check2']
DETAIL_SHEET_PREFIX     = RULES['detail_sheet_prefix']

BI_UNIT_IS_FEN = False

print("="*60)
print(f"项目6-系统导出数据核验 启动")
print(f"当前时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print(f"输入文件夹：{INPUT_FOLDER}")
print(f"输出文件：{OUTPUT_FILE}")
print("="*60)
print("")

# ==================== 列名重命名 ====================
def rename_columns(df):
    return df.rename(columns=FIELD_MAP)

# ==================== 加载组织层级目录 ====================
def load_valid_second_arch_labels():
    file_path = SECOND_ARCH_FILE
    if not os.path.exists(file_path):
        print(f"警告：未找到组织层级目录文件：{file_path}，跳过校验")
        return set()
    try:
        df = pd.read_excel(file_path, dtype=str)
        df.columns = [str(c).strip() for c in df.columns]
        target_col = None
        for col in df.columns:
            if '标签' in col or '部门' in col:
                target_col = col
                break
        if target_col is None:
            target_col = df.columns[0]
        labels = set(df[target_col].dropna().astype(str).str.strip())
        print(f"已加载组织层级目录，共 {len(labels)} 个有效标签")
        return labels
    except Exception as e:
        print(f"加载组织层级目录失败：{e}，跳过校验")
        return set()

VALID_SECOND_ARCH_LABELS = load_valid_second_arch_labels()

def validate_second_arch(df, valid_labels):
    errors = []
    if COL_ARCH2 not in df.columns or COL_ARCH1 not in df.columns:
        return df, pd.DataFrame()
    df[COL_ARCH2] = df[COL_ARCH2].replace(ARCH_REPLACE_MAP)
    for idx, row in df.iterrows():
        second = row.get(COL_ARCH2)
        first = row.get(COL_ARCH1)
        reasons = []
        if pd.notna(second) and str(second).strip().endswith('分公司'):
            if pd.isna(first) or str(first).strip() != '分公司':
                reasons.append(f"二级部门为'{second}'，但一级部门不是'分公司'（实际：{first}）")
        if pd.notna(second) and str(second).strip() not in valid_labels:
            reasons.append(f"二级部门'{second}'不在组织层级目录中")
        if reasons:
            errors.append({
                '行索引': idx,
                '组织层级1': first,
                '组织层级2': second,
                '错误原因': '；'.join(reasons)
            })
    errors_df = pd.DataFrame(errors)
    if not errors_df.empty:
        print(f"   [警告] 发现 {len(errors_df)} 行组织层级校验错误。")
    return df, errors_df

# ==================== 辅助函数 ====================
def find_input_file(folder):
    exts = (".xlsx", ".xls")
    files = [f for f in os.listdir(folder) if f.lower().endswith(exts) and not f.startswith('~$')]
    if not files:
        return None
    if len(files) == 1:
        return os.path.join(folder, files[0])
    else:
        print("发现多个 Excel 文件，请选择要处理的文件：")
        for i, f in enumerate(files, 1):
            print(f"  {i}. {f}")
        while True:
            try:
                choice = input("请输入序号（1-{}）：".format(len(files)))
                idx = int(choice) - 1
                if 0 <= idx < len(files):
                    return os.path.join(folder, files[idx])
                else:
                    print("序号无效，请重新输入。")
            except ValueError:
                print("请输入数字序号。")

def load_match_table():
    try:
        df = pd.read_excel(MATCH_FILE, sheet_name=MATCH_SHEET, dtype=str)
        df = rename_columns(df)
        df.columns = [col.strip() for col in df.columns]
        if COL_BIZ_CAT2 in df.columns:
            return df[[COL_BIZ_CAT2, COL_CHANNEL, COL_BIZ_CAT1]].drop_duplicates(subset=[COL_BIZ_CAT2])
        return None
    except Exception as e:
        print(f"加载匹配表失败：{e}")
        return None

def find_income_cost_columns(col_list):
    income_patterns = ['补贴金额', '收入', '含税收入', '奖励']
    cost_patterns = ['支出金额', '成本', '含税成本', '结算']
    income_col = None
    cost_col = None
    for col in col_list:
        col_lower = col.lower()
        if income_col is None:
            for pattern in income_patterns:
                if pattern in col_lower:
                    income_col = col
                    break
        if cost_col is None:
            for pattern in cost_patterns:
                if pattern in col_lower:
                    cost_col = col
                    break
        if income_col is not None and cost_col is not None:
            break
    return income_col, cost_col

def find_category_column(col_list):
    for col in col_list:
        col_lower = col.lower()
        if col_lower in [COL_CATEGORY, '一级分类', '类型']:
            return col
    for col in col_list:
        if COL_CATEGORY in col:
            return col
    return None

def classify_row(row):
    second_cat = row.get(COL_BIZ_CAT2, '')
    first_cat  = row.get(COL_BIZ_CAT1, '')
    channel    = row.get(COL_CHANNEL, '')
    project_type = row.get(COL_PROJECT_TYPE, '')
    project_name = row.get(COL_PROJECT, '')

    # 直接映射
    for key, val in CLASSIFY_DIRECT_MAP.items():
        if second_cat == key:
            return val
    # 普通分润拆分
    if CLASSIFY_SPLIT_KEYWORD in str(second_cat):
        combined = f"{project_type}{project_name}"
        if any(kw in combined for kw in CLASSIFY_KEYWORDS):
            return CLASSIFY_TARGET_CATEGORY
        return CLASSIFY_SPLIT_KEYWORD
    # 一级分类映射
    if str(first_cat) in SPECIAL_CATS:
        return first_cat
    # 通道合同
    if channel == '通道合同':
        return '通道合同收入'
    # 支付宝/微信归类到项目
    if channel in ['支付宝', '微信']:
        return CLASSIFY_TARGET_CATEGORY
    return CLASSIFY_TARGET_CATEGORY

def map_secondary_category_for_check(df):
    def get_mapped(row):
        second = row.get(COL_BIZ_CAT2, '')
        proj_name = row.get(COL_PROJECT, '')
        for src, kws in CLASSIFY_SPECIAL_MAP.items():
            if second == src:
                for kw in kws:
                    if kw in str(proj_name):
                        return f"{src}-{kw}"
        return second
    df['匹配业务小类'] = df.apply(get_mapped, axis=1)
    return df

def merge_sheets(input_path):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] 开始读取明细表...")
    xl = pd.ExcelFile(input_path)
    sheet_names = xl.sheet_names
    detail_sheets = [s for s in sheet_names if s.startswith(DETAIL_SHEET_PREFIX)]
    detail_sheets.sort(key=lambda x: int(x.split('_')[1]) if '_' in x else 0)
    if not detail_sheets:
        raise ValueError("未找到任何明细表")
    print(f"发现 {len(detail_sheets)} 个明细表")
    dfs = []
    header_cols = None
    for i, sheet in enumerate(detail_sheets):
        print(f"  读取第 {i+1}/{len(detail_sheets)} 个表: {sheet}")
        if i == 0:
            df = pd.read_excel(xl, sheet_name=sheet, dtype=str, keep_default_na=False)
            df = rename_columns(df)   # 关键：读入后脱敏
            header_cols = [str(col).strip().replace('\n', ' ') for col in df.columns]
            df.columns = header_cols
        else:
            df = pd.read_excel(xl, sheet_name=sheet, header=None, dtype=str, keep_default_na=False)
            if len(df.columns) == len(header_cols):
                df.columns = header_cols
            else:
                if len(df.columns) > len(header_cols):
                    df = df.iloc[:, :len(header_cols)]
                else:
                    for _ in range(len(header_cols) - len(df.columns)):
                        df[f"_临时列_{_}"] = pd.NA
                df.columns = header_cols
        dfs.append(df)
        print(f"    本表行数: {len(df)}")
    combined = pd.concat(dfs, ignore_index=True, sort=False)
    print(f"合并后总行数: {len(combined)}")
    del dfs, xl
    gc.collect()
    return combined

def normalize_amounts(df, cols):
    for col in cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce')
    return df

def load_check_file_transaction(file_path, sheet_name):
    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name, dtype=str)
        df = rename_columns(df)
        df.columns = [col.strip() for col in df.columns]
        if COL_BIZ_CAT2 not in df.columns:
            return None
        trans_col = None
        amount_col = None
        for col in df.columns:
            if '交易笔数' in col or '笔数' in col:
                trans_col = col
                break
        for col in df.columns:
            if '交易金额' in col or '金额' in col:
                amount_col = col
                break
        if trans_col is None or amount_col is None:
            return None
        df = df[[COL_BIZ_CAT2, trans_col, amount_col]]
        df.rename(columns={trans_col: COL_TXN_COUNT, amount_col: COL_TXN_AMT}, inplace=True)
        return df
    except Exception as e:
        print(f"  加载失败：{e}")
        return None

def load_check_file_profit(file_path, sheet_name, strict=False):
    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name, dtype=str)
        df = rename_columns(df)
        df.columns = [col.strip() for col in df.columns]
        col_list = df.columns.tolist()
        if strict:
            required = [COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST]
            for col in required:
                if col not in df.columns:
                    strict = False
                    break
            if strict:
                return df[required]
        cat_col = find_category_column(col_list)
        income_col, cost_col = find_income_cost_columns(col_list)
        if income_col is None or cost_col is None:
            return None
        has_second = COL_BIZ_CAT2 in df.columns
        if has_second:
            if cat_col is not None and cat_col != COL_BIZ_CAT2:
                df = df[[COL_BIZ_CAT2, cat_col, income_col, cost_col]]
                df.rename(columns={cat_col: COL_CATEGORY}, inplace=True)
            else:
                df = df[[COL_BIZ_CAT2, income_col, cost_col]]
                df[COL_CATEGORY] = df[COL_BIZ_CAT2]
        else:
            if cat_col is not None:
                df = df[[cat_col, income_col, cost_col]]
                df.rename(columns={cat_col: COL_CATEGORY}, inplace=True)
                df[COL_BIZ_CAT2] = df[COL_CATEGORY]
            else:
                return None
        df.rename(columns={income_col: COL_REWARD_INCOME, cost_col: COL_SETTLE_COST}, inplace=True)
        return df
    except Exception as e:
        print(f"  加载收入成本数据失败：{e}")
        return None

def compare_and_highlight(df_bi, df_offline, key_cols, value_cols, sheet_name, writer, title):
    for key in key_cols:
        if key not in df_bi.columns:
            raise KeyError(f"BI 缺少列: {key}")
        if key not in df_offline.columns:
            raise KeyError(f"线下 缺少列: {key}")
    for col in value_cols:
        if col in df_bi.columns:
            df_bi[col] = pd.to_numeric(df_bi[col], errors='coerce')
        if col in df_offline.columns:
            df_offline[col] = pd.to_numeric(df_offline[col], errors='coerce')
    bi_renamed = df_bi.rename(columns={col: f"BI_{col}" for col in df_bi.columns if col not in key_cols})
    off_renamed = df_offline.rename(columns={col: f"线下_{col}" for col in df_offline.columns if col not in key_cols})
    merged = pd.merge(bi_renamed, off_renamed, on=key_cols, how='outer')
    for col in value_cols:
        bi_col = f"BI_{col}"
        off_col = f"线下_{col}"
        diff_col = f"{col}_差异"
        if bi_col in merged.columns and off_col in merged.columns:
            merged[diff_col] = merged[bi_col].fillna(0) - merged[off_col].fillna(0)
            merged.loc[merged[bi_col].isna() & merged[off_col].isna(), diff_col] = np.nan
        elif bi_col in merged.columns:
            merged[diff_col] = merged[bi_col]
        elif off_col in merged.columns:
            merged[diff_col] = -merged[off_col]
        else:
            merged[diff_col] = np.nan
    for col in merged.columns:
        if merged[col].dtype in ['float64', 'float32']:
            merged[col] = merged[col].round(2)
    merged.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)
    workbook = writer.book
    worksheet = writer.sheets[sheet_name]
    worksheet.cell(row=1, column=1, value=title)
    worksheet.cell(row=1, column=1).font = Font(bold=True, size=12)
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    for col_idx, col_name in enumerate(merged.columns, start=1):
        if "_差异" in col_name:
            for row_idx in range(2, len(merged)+2):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if cell.value is not None and isinstance(cell.value, (int, float)) and abs(cell.value) > 0.01:
                    cell.fill = red_fill
    for col in worksheet.columns:
        column_letter = col[0].column_letter
        worksheet.column_dimensions[column_letter].width = 18

# ==================== 主函数 ====================
def main():
    try:
        print(f"[{datetime.now().strftime('%H:%M:%S')}] 开始主流程...")
        input_file = find_input_file(INPUT_FOLDER)
        if input_file is None:
            print(f"错误：未找到 Excel 文件。")
            input("按回车键退出...")
            return
        print(f"处理的文件：{input_file}")

        try:
            df = merge_sheets(input_file)
        except Exception as e:
            print(f"合并明细表失败：{e}")
            import traceback
            traceback.print_exc()
            input("按回车键退出...")
            return

        print(f"[{datetime.now().strftime('%H:%M:%S')}] 正在计算分类...")
        df[COL_CATEGORY] = df.apply(classify_row, axis=1)
        df = map_secondary_category_for_check(df)
        amount_cols = [COL_TXN_COUNT, COL_TXN_AMT, COL_REWARD_INCOME, COL_SETTLE_COST]
        df = normalize_amounts(df, amount_cols)

        # ===== 组织层级校验 =====
        print(f"[{datetime.now().strftime('%H:%M:%S')}] 正在校验组织层级...")
        df, errors_df = validate_second_arch(df, VALID_SECOND_ARCH_LABELS)

        match_df = load_match_table()
        if match_df is not None:
            df_check = df[[COL_BIZ_CAT2, COL_CHANNEL, COL_BIZ_CAT1]].drop_duplicates(subset=[COL_BIZ_CAT2])
            merged_match = pd.merge(df_check, match_df, on=COL_BIZ_CAT2, how='outer', suffixes=('_BI', '_匹配'))
            merged_match['渠道一致'] = merged_match[f'{COL_CHANNEL}_BI'] == merged_match[f'{COL_CHANNEL}_匹配']
            merged_match['业务大类一致'] = merged_match[f'{COL_BIZ_CAT1}_BI'] == merged_match[f'{COL_BIZ_CAT1}_匹配']
            result = merged_match[[COL_BIZ_CAT2, f'{COL_CHANNEL}_BI', f'{COL_CHANNEL}_匹配', '渠道一致',
                                   f'{COL_BIZ_CAT1}_BI', f'{COL_BIZ_CAT1}_匹配', '业务大类一致']]
            result.columns = [COL_BIZ_CAT2, 'BI渠道', '匹配表渠道', '渠道一致',
                              'BI业务大类', '匹配表业务大类', '业务大类一致']
            try:
                result.to_excel(MATCH_OUTPUT, index=False)
                print(f"分类校验汇总已保存至：{MATCH_OUTPUT}")
            except PermissionError:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                new_path = MATCH_OUTPUT.replace(".xlsx", f"_{timestamp}.xlsx")
                result.to_excel(new_path, index=False)
                print(f"分类校验汇总已另存至：{new_path}")

        # 处理文件占用
        output_path = OUTPUT_FILE
        if os.path.exists(output_path):
            try:
                os.remove(output_path)
                print(f"已删除旧文件：{output_path}")
            except PermissionError:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                new_path = output_path.replace(".xlsx", f"_{timestamp}.xlsx")
                print(f"原文件被占用，将另存为：{new_path}")
                output_path = new_path

        print(f"[{datetime.now().strftime('%H:%M:%S')}] 开始生成核对结果...")
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            ordinary = df[df[COL_CATEGORY] == CLASSIFY_SPLIT_KEYWORD]
            if not ordinary.empty:
                bi_trans = ordinary.groupby('匹配业务小类').agg({
                    COL_TXN_COUNT: 'sum', COL_TXN_AMT: 'sum'
                }).reset_index().rename(columns={'匹配业务小类': COL_BIZ_CAT2})
                offline_trans = load_check_file_transaction(CHECK_FILE1, SHEET_CHECK1_TRANS)
                if offline_trans is not None:
                    compare_and_highlight(bi_trans, offline_trans,
                                          key_cols=[COL_BIZ_CAT2],
                                          value_cols=[COL_TXN_COUNT, COL_TXN_AMT],
                                          sheet_name='普通分润-笔数和交易',
                                          writer=writer, title='普通分润-笔数和交易核对')
                bi_profit = ordinary.groupby('匹配业务小类').agg({
                    COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
                }).reset_index().rename(columns={'匹配业务小类': COL_BIZ_CAT2})
                offline_profit = load_check_file_profit(CHECK_FILE1, SHEET_CHECK1_PROFIT, strict=False)
                if offline_profit is not None:
                    compare_and_highlight(bi_profit, offline_profit,
                                          key_cols=[COL_BIZ_CAT2],
                                          value_cols=[COL_REWARD_INCOME, COL_SETTLE_COST],
                                          sheet_name='普通分润-分润核对',
                                          writer=writer, title='普通分润-分润核对')

            special_df = df[df[COL_CATEGORY].isin(SPECIAL_CATS)]
            if not special_df.empty:
                bi_special = special_df.groupby(COL_CATEGORY).agg({
                    COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
                }).reset_index()
                offline_profit2 = load_check_file_profit(CHECK_FILE2, SHEET_CHECK2, strict=False)
                if offline_profit2 is not None:
                    offline_special = offline_profit2[offline_profit2[COL_CATEGORY].isin(SPECIAL_CATS)]
                    if not offline_special.empty:
                        offline_special = offline_special.groupby(COL_CATEGORY).agg({
                            COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
                        }).reset_index()
                        compare_and_highlight(bi_special, offline_special,
                                              key_cols=[COL_CATEGORY],
                                              value_cols=[COL_REWARD_INCOME, COL_SETTLE_COST],
                                              sheet_name='其他类型核对',
                                              writer=writer, title='其他类型核对')

            project_df = df[df[COL_CATEGORY] == CLASSIFY_TARGET_CATEGORY]
            if not project_df.empty:
                bi_project = project_df.groupby('匹配业务小类').agg({
                    COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
                }).reset_index().rename(columns={'匹配业务小类': COL_BIZ_CAT2})
                offline_profit2 = load_check_file_profit(CHECK_FILE2, SHEET_CHECK2, strict=False)
                if offline_profit2 is not None:
                    offline_project = offline_profit2[offline_profit2[COL_CATEGORY] == CLASSIFY_TARGET_CATEGORY]
                    if not offline_project.empty:
                        offline_project = offline_project[[COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST]]
                        compare_and_highlight(bi_project, offline_project,
                                              key_cols=[COL_BIZ_CAT2],
                                              value_cols=[COL_REWARD_INCOME, COL_SETTLE_COST],
                                              sheet_name='项目分润核对',
                                              writer=writer, title='项目分润核对')

            # ===== 组织层级校验错误 =====
            if not errors_df.empty:
                errors_df.to_excel(writer, sheet_name='组织层级校验错误', index=False)
                print(f"  ⚠️ 组织层级校验错误 {len(errors_df)} 条已写入核验结果。")

        print(f"[{datetime.now().strftime('%H:%M:%S')}] 核验完成：{output_path}")

    except Exception as e:
        print(f"\n[{datetime.now().strftime('%H:%M:%S')}] 发生错误：{e}")
        import traceback
        traceback.print_exc()
    finally:
        input("\n按回车键退出...")

if __name__ == "__main__":
    main()