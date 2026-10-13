import os
import sys
import gc
import json
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime
import time

# ==================== 全局错误捕获 ====================
def global_exception_handler(exc_type, exc_value, exc_traceback):
    print("\n" + "="*60)
    print("发生未捕获的异常：")
    print(f"类型：{exc_type}")
    print(f"值：{exc_value}")
    import traceback
    traceback.print_tb(exc_traceback)
    print("="*60)
    input("按回车键退出...")
    sys.exit(1)

sys.excepthook = global_exception_handler

# ==================== 敏感信息配置区（外发时替换为占位符）====================
PATHS = {
    'input_folder':     r'F:\...\project3_pre_bi_validation.py',
    'field_map_file':   r'F:\...\field_map.json',
    'match_file':       r'F:\...\活动名称对应一二级分类.xlsx',
    'check_file1':      r'F:\...\普通分润核对.xlsx',
    'check_file2':      r'F:\...\收入成本校验表--python校验.xlsx',
    'output_file':      r'F:\...\核验结果.xlsx',
    'log_file':         r'F:\...\处理日志.txt',
    'second_arch_file': r'F:\...\分公司目录-新二级部门标签.xlsx',
}

RULES = {
    'arch_replace_map': {
        '【旧分公司A】': '【新分公司A】',
        '【旧分公司B】': '【新分公司B】',
        '【旧分公司C】': '【新分公司C】',
        '【旧总部】':   '【新总部】',
    },
    'special_manager_id':  0,
    'special_second_arch': '【特殊分公司名称】',
    'exempt_category_for_agent':  '【豁免分类编码】',
    'invoice_exempt_category':    '【开票豁免分类编码】',
    'allowed_arch_exceptions': ['【例外部门1】', '【例外部门2】'],

    'match_sheet_name': '名称对应活动类型',
    'log_sheet_name_check1_trans': '笔数和交易',
    'log_sheet_name_check1_profit': '分润核对',
    'log_sheet_name_check2': 'Sheet1',

    'classify_keywords': [
        '【关键词1】', '【关键词2】', '【关键词3】', '【关键词4】',
    ],
    'classify_target_category': '项目',
    'classify_direct_map': {
        '【特殊二级分类1】': '【对应分类1】',
        '【特殊二级分类2】': '【对应分类2】',
    },
    'classify_split_keyword': '普通分润',
    'special_cats': [
        '【特殊分类1】', '【特殊分类2】', '【特殊分类3】',
        '【特殊分类4】', '【特殊分类5】', '【特殊分类6】',
        '【特殊分类7】', '【特殊分类8】', '【特殊分类9】',
    ],
    'classify_special_map': {
        '【映射源分类】': ['【关键词1】', '【关键词2】'],
    },

    'max_file_size_mb': 200,
}

# 脱敏后的字段名常量
COL_MERCHANT_NO   = '商户编号'
COL_MERCHANT_NAME = '商户名'
COL_EMP_ID        = '员工编号'
COL_EMP_NAME      = '员工姓名'
COL_AGENT_ID      = '渠道编号'
COL_AGENT_NAME    = '渠道名称'
COL_ARCH1         = '组织层级1'
COL_ARCH2         = '组织层级2'
COL_ARCH3         = '组织层级3'
COL_ARCH4         = '组织层级4'
COL_OPS_ID        = '运营经理编号'
COL_OPS_NAME      = '运营经理姓名'
COL_INDUSTRY      = '行业分类'
COL_SALES_ID      = '销售员编号'
COL_SALES_NAME    = '销售员姓名'
COL_TXN_COUNT     = '有效交易笔数'
COL_TXN_AMT       = '有效交易金额'
COL_REWARD        = '补贴金额'
COL_REWARD_BASE   = '补贴/基数'
COL_RATIO         = '分润比例'
COL_SETTLE        = '支出金额'
COL_PROFIT        = '净收益'
COL_CHANNEL       = '支付渠道'
COL_BIZ_CAT1      = '业务大类'
COL_BIZ_CAT2      = '业务小类'
COL_BILL_PERIOD   = '账期'
COL_BI_PERIOD     = '报表期'
COL_TABLE_NAME    = '数据表'
COL_INVOICE       = '发票抬头'
COL_PROJECT       = '项目名'
COL_PROJECT_TYPE  = '项目类型'
COL_PROJECT_CAT   = '项目类别'
COL_SETTLE_METHOD = '结算方式'
COL_AGENT_SETTLE  = '渠道结算标记'
COL_EMP_SETTLE    = '员工结算标记'
COL_EMP_D0        = '员工即时分润'
COL_AGENT_D0      = '渠道即时分润'

# 收入/成本列名（脱敏后）
COL_REWARD_INCOME = '补贴金额（收入）'
COL_SETTLE_COST   = '支出金额（成本）'
COL_CATEGORY      = '分类'
# ==========================================================================

with open(PATHS['field_map_file'], 'r', encoding='utf-8') as f:
    FIELD_MAP = json.load(f)

INPUT_FOLDER     = PATHS['input_folder']
MATCH_FILE       = PATHS['match_file']
MATCH_SHEET      = RULES['match_sheet_name']
CHECK_FILE1      = PATHS['check_file1']
CHECK_FILE2      = PATHS['check_file2']
OUTPUT_FILE      = PATHS['output_file']
LOG_FILE         = PATHS['log_file']
SECOND_ARCH_FILE = PATHS['second_arch_file']

SPECIAL_MANAGER_ID        = RULES['special_manager_id']
SPECIAL_SECOND_ARCH       = RULES['special_second_arch']
EXEMPT_CATEGORY_FOR_AGENT = RULES['exempt_category_for_agent']
INVOICE_EXEMPT_CATEGORY   = RULES['invoice_exempt_category']
ALLOWED_ARCH_EXCEPTIONS   = RULES['allowed_arch_exceptions']
ARCH_REPLACE_MAP          = RULES['arch_replace_map']
SPECIAL_CATS              = RULES['special_cats']
MAX_FILE_SIZE_MB          = RULES['max_file_size_mb']
CLASSIFY_KEYWORDS         = RULES['classify_keywords']
CLASSIFY_TARGET_CATEGORY  = RULES['classify_target_category']
CLASSIFY_DIRECT_MAP       = RULES['classify_direct_map']
CLASSIFY_SPLIT_KEYWORD    = RULES['classify_split_keyword']
CLASSIFY_SPECIAL_MAP      = RULES['classify_special_map']
SHEET_CHECK1_TRANS        = RULES['log_sheet_name_check1_trans']
SHEET_CHECK1_PROFIT       = RULES['log_sheet_name_check1_profit']
SHEET_CHECK2              = RULES['log_sheet_name_check2']

BI_UNIT_IS_FEN = True

print("="*60)
print(f"项目3-导入前数据核验")
print(f"当前时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print(f"输入文件夹：{INPUT_FOLDER}")
print(f"输出文件：{OUTPUT_FILE}")
print("="*60)
print("")

# ==================== 列名重命名 ====================
def rename_columns(df):
    return df.rename(columns=FIELD_MAP)

# ==================== 加载组织层级目录 ====================
def load_valid_second_arch_labels():
    file_path = SECOND_ARCH_FILE
    if not os.path.exists(file_path):
        print(f"警告：未找到组织层级目录文件：{file_path}，跳过校验")
        return set()
    try:
        df = pd.read_excel(file_path, dtype=str)
        df.columns = [str(c).strip() for c in df.columns]
        target_col = None
        for col in df.columns:
            if '标签' in col or '部门' in col:
                target_col = col
                break
        if target_col is None:
            target_col = df.columns[0]
        labels = set(df[target_col].dropna().astype(str).str.strip())
        print(f"已加载组织层级目录，共 {len(labels)} 个有效标签")
        return labels
    except Exception as e:
        print(f"加载组织层级目录失败：{e}，跳过校验")
        return set()

VALID_SECOND_ARCH_LABELS = load_valid_second_arch_labels()

def validate_second_arch(df, valid_labels):
    errors = []
    if COL_ARCH2 not in df.columns or COL_ARCH1 not in df.columns:
        return df, pd.DataFrame()
    df[COL_ARCH2] = df[COL_ARCH2].replace(ARCH_REPLACE_MAP)
    for idx, row in df.iterrows():
        second = row.get(COL_ARCH2)
        first = row.get(COL_ARCH1)
        reasons = []
        if pd.notna(second) and str(second).strip().endswith('分公司'):
            if pd.isna(first) or str(first).strip() != '分公司':
                reasons.append(f"二级部门为'{second}'，但一级部门不是'分公司'（实际：{first}）")
        if pd.notna(second) and str(second).strip() not in valid_labels:
            reasons.append(f"二级部门'{second}'不在组织层级目录中")
        if reasons:
            errors.append({
                '行索引': idx,
                '组织层级1': first,
                '组织层级2': second,
                '错误原因': '；'.join(reasons)
            })
    errors_df = pd.DataFrame(errors)
    if not errors_df.empty:
        print(f"   [警告] 发现 {len(errors_df)} 行组织层级校验错误。")
    return df, errors_df

# ==================== 辅助函数 ====================
def classify_row(row):
    second_cat = row.get(COL_BIZ_CAT2, '')
    first_cat  = row.get(COL_BIZ_CAT1, '')
    channel    = row.get(COL_CHANNEL, '')
    project_type = row.get(COL_PROJECT_TYPE, '')
    project_name = row.get(COL_PROJECT, '')

    # 直接映射
    for key, val in CLASSIFY_DIRECT_MAP.items():
        if second_cat == key:
            return val
    # 普通分润拆分
    if CLASSIFY_SPLIT_KEYWORD in str(second_cat):
        combined = f"{project_type}{project_name}"
        if any(kw in combined for kw in CLASSIFY_KEYWORDS):
            return CLASSIFY_TARGET_CATEGORY
        return CLASSIFY_SPLIT_KEYWORD
    # 一级分类映射
    if str(first_cat) in SPECIAL_CATS:
        return first_cat
    # 通道合同
    if channel == '通道合同':
        return '通道合同收入'
    # 支付宝/微信归类到项目
    if channel in ['支付宝', '微信']:
        return CLASSIFY_TARGET_CATEGORY
    return CLASSIFY_TARGET_CATEGORY

def map_secondary_category_for_check(df):
    def get_mapped(row):
        second = row.get(COL_BIZ_CAT2, '')
        proj_name = row.get(COL_PROJECT, '')
        for src, kws in CLASSIFY_SPECIAL_MAP.items():
            if second == src:
                for kw in kws:
                    if kw in str(proj_name):
                        return f"{src}-{kw}"
        return second
    df['匹配业务小类'] = df.apply(get_mapped, axis=1)
    return df

def find_income_cost_columns(col_list):
    income_patterns = ['补贴金额', '收入', '含税收入', '奖励']
    cost_patterns = ['支出金额', '成本', '含税成本', '结算']
    income_col = None
    cost_col = None
    for col in col_list:
        col_lower = col.lower()
        if income_col is None:
            for pattern in income_patterns:
                if pattern in col_lower:
                    income_col = col
                    break
        if cost_col is None:
            for pattern in cost_patterns:
                if pattern in col_lower:
                    cost_col = col
                    break
        if income_col is not None and cost_col is not None:
            break
    return income_col, cost_col

def find_category_column(col_list):
    for col in col_list:
        col_lower = col.lower()
        if col_lower in [COL_CATEGORY, '一级分类', '类型']:
            return col
    for col in col_list:
        if COL_CATEGORY in col:
            return col
    return None

def merge_sheets_from_file(file_path, log_handle=None):
    file_name = os.path.basename(file_path)
    try:
        size_mb = os.path.getsize(file_path) / (1024 * 1024)
        if size_mb > MAX_FILE_SIZE_MB:
            msg = f"  跳过文件 {file_name}：文件过大 ({size_mb:.1f} MB)"
            print(msg)
            if log_handle:
                log_handle.write(msg + "\n")
            return None
    except Exception as e:
        msg = f"  跳过文件 {file_name}：无法获取文件大小（{e}）"
        print(msg)
        if log_handle:
            log_handle.write(msg + "\n")
        return None
    print(f"  读取文件：{file_name}")
    try:
        df = pd.read_excel(file_path, sheet_name=0, dtype=str, keep_default_na=False, engine='openpyxl')
        df = rename_columns(df)   # 关键：读入后脱敏
        df.columns = [col.strip() for col in df.columns]
        print(f"    本文件读取 {len(df)} 行")
        return df
    except Exception as e:
        msg = f"  跳过文件 {file_name}：读取失败（{e}）"
        print(msg)
        if log_handle:
            log_handle.write(msg + "\n")
        return None

def standardize_columns(df, log_handle=None):
    mapping = {
        COL_REWARD_INCOME:  ['补贴金额', '含税收入', '收入', '奖励金额(收入)'],
        COL_SETTLE_COST:    ['支出金额', '含税成本', '成本', '结算金额(成本)'],
        COL_TXN_AMT:        ['交易金额', '有效金额', '交易金额合计'],
        COL_TXN_COUNT:      ['交易笔数', '有效笔数', '笔数'],
        COL_BIZ_CAT2:       ['业务二级分类', '二级分类', '分类名称'],
        COL_BIZ_CAT1:       ['业务一级分类', '一级分类', '大类'],
        COL_CHANNEL:        ['通道', '支付通道'],
        COL_PROJECT_TYPE:   ['项目类型', '类型'],
        COL_PROJECT:        ['项目名称', '活动名称'],
        COL_INVOICE:        ['开票名称', '发票抬头'],
        COL_EMP_NAME:       ['业务经理名称', '经理名称'],
        COL_AGENT_NAME:     ['代理商名称', '代理名称'],
        COL_BILL_PERIOD:    ['账单期间', '期间'],
        COL_ARCH1:          ['公司1级行政架构', '一级部门', '一级行政架构'],
        COL_ARCH2:          ['公司2级行政架构', '二级部门', '二级行政架构'],
    }
    for std_col, candidates in mapping.items():
        need_fill = False
        if std_col not in df.columns:
            need_fill = True
        else:
            col_series = df[std_col]
            if col_series.isna().all():
                need_fill = True
            else:
                non_null = col_series.dropna()
                if len(non_null) > 0 and (non_null.astype(str).str.strip() == '').all():
                    need_fill = True
        if need_fill:
            for cand in candidates:
                if cand in df.columns:
                    df[std_col] = df[cand]
                    print(f"    列名映射：'{cand}' -> '{std_col}'")
                    if log_handle:
                        log_handle.write(f"列名映射：'{cand}' -> '{std_col}'\n")
                    break
    for col in [COL_REWARD_INCOME, COL_SETTLE_COST, COL_TXN_AMT, COL_TXN_COUNT,
                COL_BIZ_CAT2, COL_BIZ_CAT1, COL_CHANNEL, COL_PROJECT_TYPE, COL_PROJECT,
                COL_INVOICE, COL_EMP_NAME, COL_AGENT_NAME, COL_BILL_PERIOD,
                COL_ARCH1, COL_ARCH2]:
        if col not in df.columns:
            df[col] = pd.NA
    return df

def load_match_table():
    try:
        df = pd.read_excel(MATCH_FILE, sheet_name=MATCH_SHEET, dtype=str)
        df = rename_columns(df)
        df.columns = [col.strip() for col in df.columns]
        if COL_BIZ_CAT2 in df.columns:
            return df[[COL_BIZ_CAT2, COL_CHANNEL, COL_BIZ_CAT1]].drop_duplicates(subset=[COL_BIZ_CAT2])
        return None
    except Exception as e:
        print(f"加载匹配表失败：{e}")
        return None

def load_check_file_transaction(file_path, sheet_name):
    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name, dtype=str)
        df = rename_columns(df)
        df.columns = [col.strip() for col in df.columns]
        if COL_BIZ_CAT2 not in df.columns:
            return None
        trans_col = None
        amount_col = None
        for col in df.columns:
            if '交易笔数' in col or '笔数' in col:
                trans_col = col
                break
        for col in df.columns:
            if '交易金额' in col or '金额' in col:
                amount_col = col
                break
        if trans_col is None or amount_col is None:
            return None
        df = df[[COL_BIZ_CAT2, trans_col, amount_col]]
        df.rename(columns={trans_col: COL_TXN_COUNT, amount_col: COL_TXN_AMT}, inplace=True)
        for col in [COL_TXN_COUNT, COL_TXN_AMT]:
            df[col] = pd.to_numeric(df[col], errors='coerce')
        return df
    except Exception as e:
        print(f"  加载失败：{e}")
        return None

def load_check_file_profit(file_path, sheet_name, strict=False):
    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name, dtype=str)
        df = rename_columns(df)
        df.columns = [col.strip() for col in df.columns]
        col_list = df.columns.tolist()
        if strict:
            required = [COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST]
            for col in required:
                if col not in df.columns:
                    return None
            return df[required]
        cat_col = find_category_column(col_list)
        income_col, cost_col = find_income_cost_columns(col_list)
        if income_col is None or cost_col is None:
            return None
        has_second = COL_BIZ_CAT2 in df.columns
        if has_second:
            if cat_col is not None and cat_col != COL_BIZ_CAT2:
                df = df[[COL_BIZ_CAT2, cat_col, income_col, cost_col]]
                df.rename(columns={cat_col: COL_CATEGORY}, inplace=True)
            else:
                df = df[[COL_BIZ_CAT2, income_col, cost_col]]
                df[COL_CATEGORY] = df[COL_BIZ_CAT2]
        else:
            if cat_col is not None:
                df = df[[cat_col, income_col, cost_col]]
                df.rename(columns={cat_col: COL_CATEGORY}, inplace=True)
                df[COL_BIZ_CAT2] = df[COL_CATEGORY]
            else:
                return None
        df.rename(columns={income_col: COL_REWARD_INCOME, cost_col: COL_SETTLE_COST}, inplace=True)
        for col in [COL_REWARD_INCOME, COL_SETTLE_COST]:
            df[col] = pd.to_numeric(df[col], errors='coerce')
        return df
    except Exception as e:
        print(f"  加载收入成本数据失败：{e}")
        return None

def compare_and_highlight(df_bi, df_offline, key_cols, value_cols, sheet_name, writer, title):
    for key in key_cols:
        if key not in df_bi.columns:
            raise KeyError(f"BI 缺少列: {key}")
        if key not in df_offline.columns:
            raise KeyError(f"线下 缺少列: {key}")
    for col in value_cols:
        if col in df_bi.columns:
            df_bi[col] = pd.to_numeric(df_bi[col], errors='coerce')
        if col in df_offline.columns:
            df_offline[col] = pd.to_numeric(df_offline[col], errors='coerce')
    bi_renamed = df_bi.rename(columns={col: f"BI_{col}" for col in df_bi.columns if col not in key_cols})
    off_renamed = df_offline.rename(columns={col: f"线下_{col}" for col in df_offline.columns if col not in key_cols})
    merged = pd.merge(bi_renamed, off_renamed, on=key_cols, how='outer')
    for col in value_cols:
        bi_col = f"BI_{col}"
        off_col = f"线下_{col}"
        diff_col = f"{col}_差异"
        if bi_col in merged.columns and off_col in merged.columns:
            merged[diff_col] = merged[bi_col].fillna(0) - merged[off_col].fillna(0)
            merged.loc[merged[bi_col].isna() & merged[off_col].isna(), diff_col] = np.nan
        elif bi_col in merged.columns:
            merged[diff_col] = merged[bi_col]
        elif off_col in merged.columns:
            merged[diff_col] = -merged[off_col]
        else:
            merged[diff_col] = np.nan
    for col in merged.columns:
        if merged[col].dtype in ['float64', 'float32']:
            merged[col] = merged[col].round(2)
    merged.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)
    workbook = writer.book
    worksheet = writer.sheets[sheet_name]
    worksheet.cell(row=1, column=1, value=title)
    worksheet.cell(row=1, column=1).font = Font(bold=True, size=12)
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    for col_idx, col_name in enumerate(merged.columns, start=1):
        if "_差异" in col_name:
            for row_idx in range(2, len(merged)+2):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if cell.value is not None and isinstance(cell.value, (int, float)) and abs(cell.value) > 0.01:
                    cell.fill = red_fill
    for col in worksheet.columns:
        column_letter = col[0].column_letter
        worksheet.column_dimensions[column_letter].width = 18

def format_bi_periods(periods):
    if periods is None or len(periods) == 0:
        return ""
    dt_list = []
    for p in periods:
        if pd.isna(p):
            continue
        try:
            if isinstance(p, datetime):
                dt = p
            elif isinstance(p, pd.Timestamp):
                dt = p.to_pydatetime()
            else:
                dt = pd.to_datetime(str(p)).to_pydatetime()
            if pd.isna(dt) or dt is pd.NaT:
                continue
            dt_list.append(dt)
        except:
            continue
    if not dt_list:
        return ""
    dt_list = sorted(set(dt_list))
    result = []
    start = dt_list[0]
    end = dt_list[0]
    for dt in dt_list[1:]:
        if pd.isna(dt) or dt is pd.NaT:
            continue
        next_month = (end.year, end.month + 1 if end.month < 12 else 1, 1)
        if (dt.year, dt.month) == (next_month[0], next_month[1]):
            end = dt
        else:
            if start == end:
                result.append(start.strftime("%Y/%m"))
            else:
                result.append(f"{start.strftime('%Y/%m')}-{end.strftime('%Y/%m')}")
            start = dt
            end = dt
    if start == end:
        result.append(start.strftime("%Y/%m"))
    else:
        result.append(f"{start.strftime('%Y/%m')}-{end.strftime('%Y/%m')}")
    return ",".join(result)

def concat_periods_raw(periods):
    if periods is None or len(periods) == 0:
        return ""
    unique_vals = []
    seen = set()
    for p in periods:
        if pd.isna(p):
            continue
        s = str(p).strip()
        if s and s not in seen:
            seen.add(s)
            unique_vals.append(s)
    return ",".join(unique_vals)

# ==================== 主函数 ====================
def main():
    print(f"[{datetime.now().strftime('%H:%M:%S')}] 开始处理...")

    if os.path.exists(OUTPUT_FILE):
        try:
            os.remove(OUTPUT_FILE)
            print(f"已删除旧文件：{OUTPUT_FILE}")
        except PermissionError:
            print(f"错误：文件 {OUTPUT_FILE} 被占用。")
            input("按回车键退出...")
            return

    with open(LOG_FILE, 'w', encoding='utf-8') as log:
        log.write(f"处理开始时间：{datetime.now()}\n")

        exts = (".xlsx", ".xls")
        files = [f for f in os.listdir(INPUT_FOLDER) if f.lower().endswith(exts) and not f.startswith('~$')]
        if not files:
            print(f"错误：未找到 Excel 文件。")
            input("按回车键退出...")
            return
        print(f"找到 {len(files)} 个文件")

        all_dfs = []
        for f in files:
            full_path = os.path.join(INPUT_FOLDER, f)
            df_file = merge_sheets_from_file(full_path, log_handle=log)
            if df_file is not None and not df_file.empty:
                df_file['来源文件'] = f
                all_dfs.append(df_file)
        if not all_dfs:
            print("没有读取到任何有效数据。")
            input("按回车键退出...")
            return

        df_all = pd.concat(all_dfs, ignore_index=True)
        print(f"合并所有文件后总行数：{len(df_all)}")

        print("正在进行列名标准化映射...")
        df_all = standardize_columns(df_all, log_handle=log)

        amount_cols = [COL_TXN_AMT, COL_REWARD_INCOME, COL_SETTLE_COST]
        for col in amount_cols:
            if col in df_all.columns:
                df_all[col] = pd.to_numeric(df_all[col], errors='coerce')
                if BI_UNIT_IS_FEN:
                    df_all[col] = df_all[col] / 100.0
        if COL_TXN_COUNT in df_all.columns:
            df_all[COL_TXN_COUNT] = pd.to_numeric(df_all[COL_TXN_COUNT], errors='coerce')

        print(f"[{datetime.now().strftime('%H:%M:%S')}] 正在计算分类...")
        df_all[COL_CATEGORY] = df_all.apply(classify_row, axis=1)
        df_all = map_secondary_category_for_check(df_all)

        print(f"[{datetime.now().strftime('%H:%M:%S')}] 正在校验组织层级...")
        df_all, errors_df = validate_second_arch(df_all, VALID_SECOND_ARCH_LABELS)
        if not errors_df.empty and '来源文件' in df_all.columns:
            errors_df['来源文件'] = df_all.loc[errors_df['行索引'], '来源文件'].values

        class_counts = df_all[COL_CATEGORY].value_counts()
        print("分类统计：")
        for cat, cnt in class_counts.items():
            print(f"  {cat}: {cnt} 行")

        invoice_errors = pd.DataFrame()
        if COL_INVOICE in df_all.columns and COL_EMP_NAME in df_all.columns and COL_AGENT_NAME in df_all.columns:
            def is_valid_invoice(row):
                invoice = row.get(COL_INVOICE)
                if pd.isna(invoice) or str(invoice).strip() == '':
                    return True
                if row.get(COL_BIZ_CAT2) == INVOICE_EXEMPT_CATEGORY:
                    return True
                manager = row.get(COL_EMP_NAME)
                agent = row.get(COL_AGENT_NAME)
                invoice_str = str(invoice).strip()
                manager_str = str(manager).strip() if not pd.isna(manager) else ''
                agent_str = str(agent).strip() if not pd.isna(agent) else ''
                return invoice_str == manager_str or invoice_str == agent_str
            invalid_mask = ~df_all.apply(is_valid_invoice, axis=1)
            if invalid_mask.any():
                invoice_errors = df_all[invalid_mask].copy()
                print(f"   [警告] 发现 {len(invoice_errors)} 行发票抬头异常。")

        # 分类校验汇总
        match_df = load_match_table()
        classification_summary = pd.DataFrame()
        if match_df is not None:
            df_check = df_all[[COL_BIZ_CAT2, COL_CHANNEL, COL_BIZ_CAT1]].drop_duplicates(subset=[COL_BIZ_CAT2])
            merged_match = pd.merge(df_check, match_df, on=COL_BIZ_CAT2, how='outer', suffixes=('_BI', '_匹配'))
            merged_match['渠道一致'] = merged_match[f'{COL_CHANNEL}_BI'] == merged_match[f'{COL_CHANNEL}_匹配']
            merged_match['业务大类一致'] = merged_match[f'{COL_BIZ_CAT1}_BI'] == merged_match[f'{COL_BIZ_CAT1}_匹配']
            result = merged_match[[COL_BIZ_CAT2, f'{COL_CHANNEL}_BI', f'{COL_CHANNEL}_匹配', '渠道一致',
                                   f'{COL_BIZ_CAT1}_BI', f'{COL_BIZ_CAT1}_匹配', '业务大类一致']]
            result.columns = [COL_BIZ_CAT2, 'BI渠道', '匹配表渠道', '渠道一致',
                              'BI业务大类', '匹配表业务大类', '业务大类一致']
            classification_summary = result

        # 项目分润核对
        project_df = df_all[df_all[COL_CATEGORY] == CLASSIFY_TARGET_CATEGORY]
        if not project_df.empty:
            bi_project = project_df.groupby('匹配业务小类').agg({
                COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
            }).reset_index().rename(columns={'匹配业务小类': COL_BIZ_CAT2})
            bi_periods = project_df.groupby('匹配业务小类')[COL_BI_PERIOD].apply(lambda x: format_bi_periods(x)).reset_index()
            bi_periods.rename(columns={'匹配业务小类': COL_BIZ_CAT2, COL_BI_PERIOD: '报表期汇总'}, inplace=True)
            bi_project = pd.merge(bi_project, bi_periods, on=COL_BIZ_CAT2, how='left')
        else:
            bi_project = pd.DataFrame(columns=[COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST, '报表期汇总'])

        offline_project = pd.DataFrame()
        try:
            df_offline_raw = pd.read_excel(CHECK_FILE2, sheet_name=SHEET_CHECK2, dtype=str)
            df_offline_raw = rename_columns(df_offline_raw)
            df_offline_raw.columns = [col.strip() for col in df_offline_raw.columns]
            offline_profit = load_check_file_profit(CHECK_FILE2, SHEET_CHECK2, strict=False)
            if offline_profit is not None:
                offline_profit = offline_profit[offline_profit[COL_CATEGORY] == CLASSIFY_TARGET_CATEGORY]
                if not offline_profit.empty:
                    offline_project = offline_profit.groupby(COL_BIZ_CAT2).agg({
                        COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
                    }).reset_index()
                    if COL_BILL_PERIOD in df_offline_raw.columns:
                        off_periods = df_offline_raw.groupby(COL_BIZ_CAT2)[COL_BILL_PERIOD].apply(lambda x: concat_periods_raw(x)).reset_index()
                        off_periods.rename(columns={COL_BILL_PERIOD: '线下账期汇总'}, inplace=True)
                        offline_project = pd.merge(offline_project, off_periods, on=COL_BIZ_CAT2, how='left')
                    else:
                        offline_project['线下账期汇总'] = ""
                else:
                    offline_project = pd.DataFrame(columns=[COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST, '线下账期汇总'])
            else:
                offline_project = pd.DataFrame(columns=[COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST, '线下账期汇总'])
        except Exception as e:
            print(f"加载线下项目数据失败：{e}")
            offline_project = pd.DataFrame(columns=[COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST, '线下账期汇总'])

        # 其他类型核对
        special_df = df_all[df_all[COL_CATEGORY].isin(SPECIAL_CATS)]
        bi_special_summary = pd.DataFrame()
        bi_special_detail = pd.DataFrame()
        if not special_df.empty:
            bi_special_summary = special_df.groupby(COL_CATEGORY).agg({
                COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
            }).reset_index()
            bi_periods_special = special_df.groupby(COL_CATEGORY)[COL_BI_PERIOD].apply(lambda x: format_bi_periods(x)).reset_index()
            bi_periods_special.rename(columns={COL_BI_PERIOD: '报表期汇总'}, inplace=True)
            bi_special_summary = pd.merge(bi_special_summary, bi_periods_special, on=COL_CATEGORY, how='left')
            bi_special_detail = special_df[[COL_CATEGORY, COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST, COL_BILL_PERIOD, COL_BI_PERIOD]].copy()

        offline_special_summary = pd.DataFrame()
        offline_special_detail = pd.DataFrame()
        try:
            df_offline_raw2 = pd.read_excel(CHECK_FILE2, sheet_name=SHEET_CHECK2, dtype=str)
            df_offline_raw2 = rename_columns(df_offline_raw2)
            df_offline_raw2.columns = [col.strip() for col in df_offline_raw2.columns]
            offline_profit2 = load_check_file_profit(CHECK_FILE2, SHEET_CHECK2, strict=False)
            if offline_profit2 is not None:
                offline_profit2 = offline_profit2[offline_profit2[COL_CATEGORY].isin(SPECIAL_CATS)]
                if not offline_profit2.empty:
                    offline_special_summary = offline_profit2.groupby(COL_CATEGORY).agg({
                        COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
                    }).reset_index()
                    if COL_BILL_PERIOD in df_offline_raw2.columns:
                        off_periods2 = df_offline_raw2.groupby(COL_CATEGORY)[COL_BILL_PERIOD].apply(lambda x: concat_periods_raw(x)).reset_index()
                        off_periods2.rename(columns={COL_BILL_PERIOD: '线下账期汇总'}, inplace=True)
                        offline_special_summary = pd.merge(offline_special_summary, off_periods2, on=COL_CATEGORY, how='left')
                    else:
                        offline_special_summary['线下账期汇总'] = ""
                    offline_special_detail = offline_profit2[[COL_CATEGORY, COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST]].copy()
                    if COL_BILL_PERIOD in df_offline_raw2.columns:
                        offline_special_detail[COL_BILL_PERIOD] = df_offline_raw2.loc[offline_special_detail.index, COL_BILL_PERIOD]
                    else:
                        offline_special_detail[COL_BILL_PERIOD] = ""
        except Exception as e:
            print(f"加载线下特殊数据失败：{e}")

        # ==================== 写入结果 ====================
        with pd.ExcelWriter(OUTPUT_FILE, engine='openpyxl') as writer:
            classification_summary.to_excel(writer, sheet_name='分类校验汇总', index=False)

            ordinary = df_all[df_all[COL_CATEGORY] == CLASSIFY_SPLIT_KEYWORD]
            if not ordinary.empty:
                bi_trans = ordinary.groupby('匹配业务小类').agg({
                    COL_TXN_COUNT: 'sum', COL_TXN_AMT: 'sum'
                }).reset_index().rename(columns={'匹配业务小类': COL_BIZ_CAT2})
                offline_trans = load_check_file_transaction(CHECK_FILE1, SHEET_CHECK1_TRANS)
                if offline_trans is not None:
                    compare_and_highlight(bi_trans, offline_trans,
                                          key_cols=[COL_BIZ_CAT2],
                                          value_cols=[COL_TXN_COUNT, COL_TXN_AMT],
                                          sheet_name='普通分润-笔数和交易',
                                          writer=writer, title='普通分润-笔数和交易核对')
                bi_profit = ordinary.groupby('匹配业务小类').agg({
                    COL_REWARD_INCOME: 'sum', COL_SETTLE_COST: 'sum'
                }).reset_index().rename(columns={'匹配业务小类': COL_BIZ_CAT2})
                offline_profit = load_check_file_profit(CHECK_FILE1, SHEET_CHECK1_PROFIT, strict=True)
                if offline_profit is not None:
                    compare_and_highlight(bi_profit, offline_profit,
                                          key_cols=[COL_BIZ_CAT2],
                                          value_cols=[COL_REWARD_INCOME, COL_SETTLE_COST],
                                          sheet_name='普通分润-分润核对',
                                          writer=writer, title='普通分润-分润核对')

            # 其他类型核对_汇总
            if not bi_special_summary.empty or not offline_special_summary.empty:
                merged_special = pd.merge(bi_special_summary, offline_special_summary,
                                          on=COL_CATEGORY, how='outer', suffixes=('_BI', '_线下'))
                for col in [COL_REWARD_INCOME, COL_SETTLE_COST]:
                    bi_col = f"{col}_BI"
                    off_col = f"{col}_线下"
                    diff_col = f"{col}_差异"
                    if bi_col in merged_special.columns and off_col in merged_special.columns:
                        merged_special[diff_col] = merged_special[bi_col].fillna(0) - merged_special[off_col].fillna(0)
                        merged_special.loc[merged_special[bi_col].isna() & merged_special[off_col].isna(), diff_col] = np.nan
                if '报表期汇总' in bi_special_summary.columns:
                    merged_special['报表期汇总'] = merged_special[COL_CATEGORY].map(bi_special_summary.set_index(COL_CATEGORY)['报表期汇总'])
                if '线下账期汇总' in offline_special_summary.columns:
                    merged_special['线下账期汇总'] = merged_special[COL_CATEGORY].map(offline_special_summary.set_index(COL_CATEGORY)['线下账期汇总'])
                merged_special = merged_special[[
                    COL_CATEGORY,
                    f'{COL_REWARD_INCOME}_BI', f'{COL_REWARD_INCOME}_线下', f'{COL_REWARD_INCOME}_差异',
                    f'{COL_SETTLE_COST}_BI', f'{COL_SETTLE_COST}_线下', f'{COL_SETTLE_COST}_差异',
                    '报表期汇总', '线下账期汇总'
                ]]
                merged_special.columns = [
                    COL_CATEGORY,
                    f'BI_{COL_REWARD_INCOME}', f'线下_{COL_REWARD_INCOME}', f'{COL_REWARD_INCOME}差异',
                    f'BI_{COL_SETTLE_COST}', f'线下_{COL_SETTLE_COST}', f'{COL_SETTLE_COST}差异',
                    '报表期汇总', '线下账期汇总'
                ]
                merged_special.to_excel(writer, sheet_name='其他类型核对_汇总', index=False)

            # 其他类型核对_明细
            if not bi_special_detail.empty or not offline_special_detail.empty:
                detail_sheet = writer.book.create_sheet("其他类型核对_明细")
                row = 1
                detail_sheet.cell(row=1, column=1, value="BI明细")
                detail_sheet.cell(row=1, column=1).font = Font(bold=True, size=12)
                row += 1
                bi_headers = [COL_CATEGORY, COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST, COL_BILL_PERIOD, COL_BI_PERIOD]
                for col_idx, h in enumerate(bi_headers, start=1):
                    detail_sheet.cell(row=row, column=col_idx, value=h)
                    detail_sheet.cell(row=row, column=col_idx).font = Font(bold=True)
                row += 1
                for _, r in bi_special_detail.iterrows():
                    for col_idx, col in enumerate(bi_headers, start=1):
                        detail_sheet.cell(row=row, column=col_idx, value=r.get(col, ''))
                    row += 1
                row += 1
                detail_sheet.cell(row=row, column=1, value="线下明细")
                detail_sheet.cell(row=row, column=1).font = Font(bold=True, size=12)
                row += 1
                off_headers = [COL_CATEGORY, COL_BIZ_CAT2, COL_REWARD_INCOME, COL_SETTLE_COST, COL_BILL_PERIOD]
                for col_idx, h in enumerate(off_headers, start=1):
                    detail_sheet.cell(row=row, column=col_idx, value=h)
                    detail_sheet.cell(row=row, column=col_idx).font = Font(bold=True)
                row += 1
                for _, r in offline_special_detail.iterrows():
                    for col_idx, col in enumerate(off_headers, start=1):
                        detail_sheet.cell(row=row, column=col_idx, value=r.get(col, ''))
                    row += 1

            # 项目分润核对
            if not bi_project.empty and not offline_project.empty:
                merged_project = pd.merge(bi_project, offline_project,
                                          on=COL_BIZ_CAT2, how='outer', suffixes=('_BI', '_线下'))
                for col in [COL_REWARD_INCOME, COL_SETTLE_COST]:
                    bi_col = f"{col}_BI"
                    off_col = f"{col}_线下"
                    diff_col = f"{col}_差异"
                    if bi_col in merged_project.columns and off_col in merged_project.columns:
                        merged_project[diff_col] = merged_project[bi_col].fillna(0) - merged_project[off_col].fillna(0)
                        merged_project.loc[merged_project[bi_col].isna() & merged_project[off_col].isna(), diff_col] = np.nan
                merged_project = merged_project[[
                    COL_BIZ_CAT2,
                    f'{COL_REWARD_INCOME}_BI', f'{COL_REWARD_INCOME}_线下', f'{COL_REWARD_INCOME}_差异',
                    f'{COL_SETTLE_COST}_BI', f'{COL_SETTLE_COST}_线下', f'{COL_SETTLE_COST}_差异',
                    '报表期汇总', '线下账期汇总'
                ]]
                merged_project.columns = [
                    COL_BIZ_CAT2,
                    f'BI_{COL_REWARD_INCOME}', f'线下_{COL_REWARD_INCOME}', f'{COL_REWARD_INCOME}差异',
                    f'BI_{COL_SETTLE_COST}', f'线下_{COL_SETTLE_COST}', f'{COL_SETTLE_COST}差异',
                    '报表期汇总', '线下账期汇总'
                ]
                merged_project.to_excel(writer, sheet_name='项目分润核对', index=False, startrow=1)
                workbook = writer.book
                worksheet = writer.sheets['项目分润核对']
                worksheet.cell(row=1, column=1, value='项目分润核对（含账期对比）')
                worksheet.cell(row=1, column=1).font = Font(bold=True, size=12)
                red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                for col_idx, col_name in enumerate(merged_project.columns, start=1):
                    if col_name in [f'{COL_REWARD_INCOME}差异', f'{COL_SETTLE_COST}差异']:
                        for row_idx in range(2, len(merged_project)+2):
                            cell = worksheet.cell(row=row_idx, column=col_idx)
                            if cell.value is not None and isinstance(cell.value, (int, float)) and abs(cell.value) > 0.01:
                                cell.fill = red_fill
            elif not bi_project.empty:
                bi_project.to_excel(writer, sheet_name='项目分润核对', index=False, startrow=1)

            if not invoice_errors.empty:
                cols = [COL_BIZ_CAT2, COL_INVOICE, COL_EMP_NAME, COL_AGENT_NAME, '来源文件']
                avail = [c for c in cols if c in invoice_errors.columns]
                invoice_errors[avail].to_excel(writer, sheet_name='发票抬头异常', index=False)

            # 组织层级校验错误
            if not errors_df.empty:
                errors_df.to_excel(writer, sheet_name='组织层级校验错误', index=False)
                print(f"  ⚠️ 组织层级校验错误 {len(errors_df)} 条已写入核验结果。")

        print(f"[{datetime.now().strftime('%H:%M:%S')}] 核验完成：{OUTPUT_FILE}")

    input("\n按回车键退出...")

if __name__ == "__main__":
    main()
